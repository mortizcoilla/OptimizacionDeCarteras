"""
scripts/md_to_docx.py
=====================

Convierte el paper markdown a un .docx con formato académico, incluyendo
figuras (placeholders [figX_nombre.png]), tablas y tipografía consistente.

Uso:
    python scripts/md_to_docx.py

Salida:
    docs/paper1_estrategias_carteras_IPSA_2010_2024.docx
"""
from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PROJECT_ROOT = Path(r"C:\Workspace\Optimizacion_de_Carteras\OptimizacionDeCarteras")
MD_PATH = Path(r"C:\Users\morti\AppData\Local\Temp\paper1.md")
FIG_DIR = PROJECT_ROOT / "docs" / "figures"
OUT_PATH = PROJECT_ROOT / "docs" / "paper1_estrategias_carteras_IPSA_2010_2024.docx"

# Colores de la paleta
COL_INK = RGBColor(0x08, 0x16, 0x30)
COL_TEAL = RGBColor(0x3B, 0x87, 0x8C)
COL_TEAL_DEEP = RGBColor(0x12, 0x53, 0x58)
COL_GREY = RGBColor(0xC2, 0xC3, 0xC5)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def set_cell_bg(cell, color_hex):
    """Set background color of a cell (RRGGBB hex without #)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'C2C3C5')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = "Calibri"
    if level == 0:
        run.font.size = Pt(22)
        run.font.bold = True
    elif level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
    else:
        run.font.size = Pt(11)
        run.font.bold = True
    run.font.color.rgb = COL_INK
    return h

def add_styled_paragraph(doc, text, italic=False, bold=False, size=11, align=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return p

def parse_inline(text):
    """Parse **bold** and *italic* markers in inline text."""
    # Return list of (text, bold, italic)
    segments = []
    # Regex to find **bold** or *italic*
    pattern = re.compile(r"(\*\*([^*]+)\*\*|\*([^*]+)\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            segments.append((text[pos:m.start()], False, False))
        if m.group(2):  # bold
            segments.append((m.group(2), True, False))
        else:  # italic
            segments.append((m.group(3), False, True))
        pos = m.end()
    if pos < len(text):
        segments.append((text[pos:], False, False))
    return segments

def add_inline_paragraph(doc, text, size=11, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    for seg_text, bold, italic in parse_inline(text):
        if not seg_text:
            continue
        run = p.add_run(seg_text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p

def add_table_from_md(doc, md_lines):
    """Parse markdown table lines and add as docx table."""
    # Header
    header = [c.strip() for c in md_lines[0].strip("|").split("|")]
    rows = []
    for line in md_lines[2:]:  # skip separator
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    # Header
    for j, txt in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(txt)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFD, 0xFC)
        set_cell_bg(cell, "125358")
    # Body
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            if j >= len(table.rows[i + 1].cells):
                continue
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run.font.size = Pt(9)
    return table

def add_figure(doc, fig_filename, caption=None, width_inches=6.0):
    fig_path = FIG_DIR / fig_filename
    if not fig_path.exists():
        # Placeholder si la figura no existe
        p = doc.add_paragraph()
        run = p.add_run(f"[Figura no encontrada: {fig_filename}]")
        run.italic = True
        run.font.color.rgb = COL_GREY
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(fig_path), width=Inches(width_inches))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        crun = cap.add_run(caption)
        crun.font.size = Pt(9)
        crun.italic = True
        crun.font.color.rgb = COL_TEAL_DEEP

# ---------------------------------------------------------------------------
# Parse markdown y generar docx
# ---------------------------------------------------------------------------

def main():
    md = MD_PATH.read_text(encoding="utf-8")
    lines = md.split("\n")

    doc = Document()

    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # Configuración de estilo Normal
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.3

    # Header del documento
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = "Ortiz C. — Estrategias de construcción de portafolios en el IPSA, 2010–2024"
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = COL_GREY
        run.italic = True

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Líneas vacías
        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith("### "):
            add_styled_heading(doc, stripped[4:], level=3)
            i += 1
            continue
        if stripped.startswith("## "):
            add_styled_heading(doc, stripped[3:], level=2)
            i += 1
            continue
        if stripped.startswith("# "):
            add_styled_heading(doc, stripped[2:], level=1)
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            add_horizontal_rule(doc)
            i += 1
            continue

        # Tablas markdown
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            # Recoger todas las líneas de la tabla
            table_lines = [line]
            j = i + 1
            while j < len(lines) and lines[j].strip().startswith("|"):
                table_lines.append(lines[j])
                j += 1
            add_table_from_md(doc, table_lines)
            i = j
            continue

        # Blockquotes (>, em uso para caption)
        if stripped.startswith("> "):
            add_styled_paragraph(doc, stripped[2:], italic=True, size=10, color=COL_TEAL_DEEP)
            i += 1
            continue

        # Figuras: [figX_nombre.png]
        m_fig = re.match(r"\[(fig\d+_[\w_]+\.png)\]", stripped)
        if m_fig:
            # Buscar caption anterior (línea que precede)
            caption = None
            if i > 0 and lines[i - 1].strip().startswith("**Figura"):
                # Tomar la línea previa completa como caption
                # Pero eso ya fue agregado como párrafo. Mejor usar un caption fijo.
                pass
            add_figure(doc, m_fig.group(1), caption=caption, width_inches=6.0)
            i += 1
            continue

        # Listas numeradas
        if re.match(r"^\d+\.\s", stripped):
            txt = re.sub(r"^\d+\.\s", "", stripped)
            doc.add_paragraph(txt, style="List Number")
            i += 1
            continue

        # Listas con guión
        if stripped.startswith("- "):
            txt = stripped[2:]
            doc.add_paragraph(txt, style="List Bullet")
            i += 1
            continue

        # Párrafo normal
        add_inline_paragraph(doc, stripped, size=11)
        i += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"OK: {OUT_PATH}")
    print(f"  Tamaño: {OUT_PATH.stat().st_size:,} bytes")

if __name__ == "__main__":
    main()
